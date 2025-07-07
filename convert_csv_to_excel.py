#!/usr/bin/env python3
"""
CSV to Excel Converter with Formatting
Converts Complete_List_of_Distinct_Codes.csv to a formatted Excel file
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter
import os
import sys

def convert_csv_to_excel(csv_path):
    """Convert CSV file to formatted Excel file"""
    
    # Check if file exists
    if not os.path.exists(csv_path):
        print(f"Error: File not found at {csv_path}")
        return False
    
    try:
        # Read the CSV file
        print(f"Reading CSV file: {csv_path}")
        df = pd.read_csv(csv_path)
        
        # Create Excel filename
        excel_path = csv_path.replace('.csv', '.xlsx')
        
        # Write to Excel
        print(f"Converting to Excel: {excel_path}")
        with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Codes', index=False)
            
            # Get the workbook and worksheet
            workbook = writer.book
            worksheet = writer.sheets['Codes']
            
            # Format headers (bold)
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # Add borders
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Apply borders to all cells with data
            for row in worksheet.iter_rows(min_row=1, max_row=len(df) + 1, 
                                         min_col=1, max_col=len(df.columns)):
                for cell in row:
                    cell.border = thin_border
            
            # Auto-fit columns
            for column_cells in worksheet.columns:
                length = max(len(str(cell.value or '')) for cell in column_cells)
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 50)
            
            # Format price column if it exists (assuming it's the second column)
            if len(df.columns) >= 2:
                for row in range(2, len(df) + 2):  # Skip header
                    cell = worksheet.cell(row=row, column=2)
                    try:
                        # Try to format as currency if it's a number
                        if pd.notna(cell.value) and isinstance(cell.value, (int, float)):
                            cell.number_format = '$#,##0.00'
                    except:
                        pass
        
        print(f"Successfully created: {excel_path}")
        print(f"Total rows converted: {len(df)}")
        print(f"Columns: {', '.join(df.columns.tolist())}")
        
        return True
        
    except Exception as e:
        print(f"Error during conversion: {str(e)}")
        return False

def main():
    # Default path
    csv_path = r'C:\Users\Happyou2\Documents\emminance\History\Complete_List_of_Distinct_Codes.csv'
    
    # Allow custom path from command line
    if len(sys.argv) > 1:
        csv_path = sys.argv[1]
    
    # Convert the file
    success = convert_csv_to_excel(csv_path)
    
    if success:
        print("\nConversion completed successfully!")
    else:
        print("\nConversion failed. Please check the error messages above.")

if __name__ == "__main__":
    main()